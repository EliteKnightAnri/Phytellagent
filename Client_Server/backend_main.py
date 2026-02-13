import json
import os
import uvicorn
import shutil
from uuid import uuid4
from typing import List, Optional
from openai import AsyncOpenAI
from fastapi import FastAPI, File, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from fastapi.responses import FileResponse
from datetime import datetime
import hashlib
from docx import Document
from pathlib import Path
from langchain_community.document_loaders import PyPDFLoader, TextLoader, Docx2txtLoader

# === 导入知识图谱模块 ===
try:
    from knowledge_graph.tool import kg_query_tool
except Exception as e:
    print(f"⚠️ KG 模块不可用: {e}")
    kg_query_tool = None


# === 导入 RAG 模块 ===
try:
    from rag_system import RAGSearch
except ImportError:
    print("❌ 警告: 未找到 rag_system.py，RAG 功能将无法使用。")
    RAGSearch = None



ALLOWED_KB_EXTS = {'.pdf', '.txt', '.md', '.docx', '.csv', '.xlsx', '.json'}
# 检查知识库上传文件的扩展名是否在白名单内
def _kb_ext_ok(filename: str) -> bool:
    _, ext = os.path.splitext(filename or "")
    return ext.lower() in ALLOWED_KB_EXTS


app = FastAPI(title="DeepSeek Agent Backend (RAG Integrated)", version="2.0")

# === 配置 API Key ===
API_KEY = "sk-f39a29f6e7e24ac687c88285fc664dd7" 

client = AsyncOpenAI(
    api_key=API_KEY,  
    base_url="https://api.deepseek.com"
)

rag_engine = None

# === 知识库 & 清单（用于“按文件开关启用/禁用”）===
KB_DIR = "source_documents"
os.makedirs(KB_DIR, exist_ok=True)

# 用 JSON 记录每个知识库文件的状态（启用/禁用、是否已索引、store_id、更新时间等）
MANIFEST_PATH = os.path.join(KB_DIR, "_kb_manifest.json")
MANIFEST_FILENAME = "_kb_manifest.json"


def _make_store_id(filename: str) -> str:
    """为知识库文件生成稳定的索引存储 ID（用于索引目录名/标识符）"""
    return hashlib.sha1(filename.encode("utf-8")).hexdigest()[:16]

def _load_manifest() -> dict:
    """从磁盘加载知识库清单；不存在或损坏时返回空结构，保证服务可用"""
    if not os.path.exists(MANIFEST_PATH):
        return {"files": {}}
    try:
        with open(MANIFEST_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        if "files" not in data or not isinstance(data["files"], dict):
            return {"files": {}}
        return data
    except Exception:
        return {"files": {}}

def _save_manifest(data: dict) -> None:
    """将清单写回磁盘（持久化 enabled/indexed/store_id 等状态）"""
    with open(MANIFEST_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def _ensure_manifest_entry(data: dict, filename: str) -> dict:
    """确保某个知识库文件在清单里有记录；若无则创建默认条目并返回该条目"""
    files = data.setdefault("files", {})
    if filename not in files:
        files[filename] = {
            "enabled": False,
            "indexed": False,
            "store_id": _make_store_id(filename),
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "updated_at": datetime.now().isoformat(timespec="seconds"),
        }
    return files[filename]

def _sync_manifest_with_disk() -> dict:
    """将 manifest 与磁盘 KB_DIR 对齐：补齐新增文件、清理已删除文件"""
    data = _load_manifest()
     # 1) 扫描磁盘：有文件就确保 manifest 里有条目
    if os.path.exists(KB_DIR):
        for f in os.listdir(KB_DIR):
            if f.startswith("."):
                continue
            _ensure_manifest_entry(data, f)

    # 2) 计算磁盘上的“有效文件集合”（排除隐藏文件和 manifest 自己）
    disk_set = set(
        x for x in os.listdir(KB_DIR)
        if (not x.startswith(".")) and x != MANIFEST_FILENAME
    ) if os.path.exists(KB_DIR) else set()

    # 3) 清理 manifest 中但磁盘已不存在的条目
    to_del = [k for k in data.get("files", {}).keys() if k not in disk_set]
    for k in to_del:
        data["files"].pop(k, None)

    _save_manifest(data)
    return data

def _bootstrap_rag_enabled_files() -> None:
    """服务启动时：根据 manifest 恢复 enabled 的知识库文件到 RAG 引擎（构建/加载索引）"""
    if not rag_engine:
        return
    data = _sync_manifest_with_disk()
    for filename, info in data.get("files", {}).items():
        if not info.get("enabled"):
            continue
        file_path = os.path.join(KB_DIR, filename)
        ok, msg = rag_engine.enable_file(filename=filename, file_path=file_path, store_id=info["store_id"])
        if ok:
            info["indexed"] = True
        else:
            # 启用失败就自动关闭，避免“显示启用但实际不可用”的状态
            info["enabled"] = False
            info["indexed"] = False
            print(f"⚠️ 启用 {filename} 失败: {msg}")
        info["updated_at"] = datetime.now().isoformat(timespec="seconds")
    _save_manifest(data)

# === 初始化 RAG 引擎（按文件开关模式）===
try:
    if RAGSearch:
        print("📚 初始化 RAG 引擎（按文件开关模式）...")
        rag_engine = RAGSearch(persist_dir="faiss_store", embedding_model="models/all-MiniLM-L6-v2")
        _bootstrap_rag_enabled_files()
        print("✅ RAG 引擎就绪！")
    else:
        print("⚠️ 未找到 rag_system 模块，RAG 模式将不可用。")
except Exception as e:
    print(f"❌ RAG 引擎初始化失败: {e}")


class FrontendChatRequest(BaseModel):
    prompt: str=""  # 用户输入的对话内容
    file_id: Optional[str] = None
    mode: Optional[str] = "general"  # "general" 或 "course_graph"
    selected_tools: Optional[List[str]] = []
    file_ids: Optional[List[str]] = None #多文件
    history: Optional[List[dict]] = []



# 对话文件上传
TEMP_DIR = "temp_uploads"
os.makedirs(TEMP_DIR, exist_ok=True)

ALLOWED_CHAT_EXTS = {'.pdf', '.txt', '.md', '.docx', '.xlsx', '.xls'}
def _chat_ext_ok(filename: str) -> bool:
    _, ext = os.path.splitext(filename or "")
    return ext.lower() in ALLOWED_CHAT_EXTS

@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
     # 后缀校验：不支持的类型直接拒绝
    if not _chat_ext_ok(file.filename):
        return {"status": "error", "message": f"不支持的文件类型：{file.filename}（支持: {', '.join(sorted(ALLOWED_CHAT_EXTS))}）"}

    file_id = str(uuid4())
    safe_name = file.filename.replace("/", "_").replace("\\", "_")
     # 保存到磁盘：文件名包含 file_id
    save_path = os.path.join(TEMP_DIR, f"{file_id}__{safe_name}")

    with open(save_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

     # 返回给前端file_id
    return {"status": "success", "file_id": file_id, "filename": file.filename}

#对话文件内容提取
def _extract_temp_text(file_path: str) -> str:
    p = Path(file_path)
    suf = p.suffix.lower()
    try:
        if suf == ".pdf":
            docs = PyPDFLoader(str(p)).load()
            return "\n".join(d.page_content for d in docs)

        if suf in [".txt", ".md"]:
            try:
                return p.read_text(encoding="utf-8")
            except:
                return p.read_text(encoding="gbk", errors="ignore")

        # xlsx
        if suf == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"【Sheet: {ws.title}】")
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if vals:
                        lines.append("\t".join(vals))
            text = "\n".join(lines)
            return text

        if suf == ".docx":
            doc = Document(str(p))
            lines = [para.text for para in doc.paragraphs if para.text and para.text.strip()]
            return "\n".join(lines)

    except Exception as e:
        print(f"[extract] failed for {file_path}: {e}")
        return ""
    return ""



#RAG检索库管理
#  1. 获取文件列表
@app.get("/files")
async def list_files():
    data = _sync_manifest_with_disk()
    items = []
    for filename, info in data.get("files", {}).items():
        items.append({
            "name": filename,
            "enabled": bool(info.get("enabled")),
            "indexed": bool(info.get("indexed")),
        })
    items.sort(key=lambda x: x["name"].lower())
    return items

# 2. 上传并索引 (覆盖之前的 mock upload)
@app.post("/upload_kb")
async def upload_knowledge_base(file: UploadFile = File(...)):
    try:
        if not _kb_ext_ok(file.filename):
            return {
                "status": "error",
                "message": f"不支持的文件类型：{file.filename}（支持: {', '.join(sorted(ALLOWED_KB_EXTS))}）"
            }

        os.makedirs(KB_DIR, exist_ok=True)

        file_path = os.path.join(KB_DIR, file.filename)
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        data = _sync_manifest_with_disk()
        info = _ensure_manifest_entry(data, file.filename)

        # 同名覆盖：索引作废（是否重建由 enabled 决定）
        info["indexed"] = False

        # 如果之前就是启用状态：只重建这个文件的索引并继续启用（不会影响其它文件）
        if rag_engine and info.get("enabled"):
            ok, msg = rag_engine.enable_file(
                filename=file.filename,
                file_path=file_path,
                store_id=info["store_id"],
            )
            if ok:
                info["indexed"] = True
            else:
                info["enabled"] = False
                info["indexed"] = False
                _save_manifest(data)
                return {"status": "error", "message": msg}

        info["updated_at"] = datetime.now().isoformat(timespec="seconds")
        _save_manifest(data)

        return {
            "status": "success",
            "message": f"上传成功：{file.filename}（已加入列表，打开开关后才会参与检索）"
        }
    except Exception as e:
        return {"status": "error", "message": str(e)}
    
# 3. 删除文件
@app.delete("/files/{filename}")
async def delete_file(filename: str):
    try:
        data = _sync_manifest_with_disk()
        info = data.get("files", {}).get(filename)
        if not info:
            return {"status": "error", "message": "文件不存在"}

        file_path = os.path.join(KB_DIR, filename)
        if os.path.exists(file_path):
            os.remove(file_path)

        if rag_engine:
            rag_engine.disable_file(filename=filename, store_id=info["store_id"], purge=True)

        data["files"].pop(filename, None)
        _save_manifest(data)

        return {"status": "success", "message": f"{filename} 已删除（已同步移除索引）"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


class ToggleKBRequest(BaseModel):
    enabled: bool



#知识库开关生效的核心接口
@app.post("/files/{filename}/toggle")
async def toggle_file(filename: str, req: ToggleKBRequest):

    # 同步 manifest 与磁盘，确保文件列表与状态是最新的
    data = _sync_manifest_with_disk()
    info = data.get("files", {}).get(filename)
    if not info:
        return {"status": "error", "message": "文件不存在"}
    
    # 二次校验
    file_path = os.path.join(KB_DIR, filename)
    if not os.path.exists(file_path):
        return {"status": "error", "message": "文件不存在（磁盘缺失）"}

    if not rag_engine:
        return {"status": "error", "message": "RAG 引擎未启动"}

    if req.enabled:
        ok, msg = rag_engine.enable_file(filename=filename, file_path=file_path, store_id=info["store_id"])
        if ok:
            info["enabled"] = True
            info["indexed"] = True
        else:
            info["enabled"] = False
            info["indexed"] = False
            _save_manifest(data)
            return {"status": "error", "message": msg}
    else:

         # 禁用：从 RAG 的启用集合移除（purge=False 表示保留索引文件以便下次快速启用）
        ok, msg = rag_engine.disable_file(filename=filename, store_id=info["store_id"], purge=False)
        info["enabled"] = False
        info["indexed"] = False

    # 更新时间戳并持久化，确保重启服务后开关状态仍能恢
    info["updated_at"] = datetime.now().isoformat(timespec="seconds")
    _save_manifest(data)
    return {"status": "success", "message": msg, "enabled": info["enabled"], "indexed": info["indexed"]}


 #返回知识库目录中的原始文件，供前端预览或下载  
@app.get("/files/download/{filename}")
async def download_file(filename: str):
    file_path = os.path.join("source_documents", filename)
    if os.path.exists(file_path):
        # filename参数让浏览器下载/打开时显示正确的文件名
        return FileResponse(file_path, filename=filename)
    return {"error": "File not found"}    




# 流式输出
@app.post("/chat")
async def chat_endpoint(request: FrontendChatRequest):
    """根据 mode 自动切换 RAG 或 通用对话"""

    async def event_generator():

        # 统一 prompt（允许为空：只发附件时给默认提示）
        user_prompt = (request.prompt or "").strip()
        if not user_prompt:
            user_prompt = "请阅读我上传的附件，并给出摘要与关键信息。"

        # 统一文件 id 列表：兼容 file_id(单个) + file_ids(多个)
        ids = []
        if getattr(request, "file_ids", None):
            ids = [x for x in (request.file_ids or []) if x]
        if not ids and getattr(request, "file_id", None):
            ids = [request.file_id]

        # 统一抽取附件文本，让 general/course_graph 都能用 
        file_context = ""
        if ids:
            file_blocks = []
            try:
                for fid in ids:
                    prefix = f"{fid}__"
                    temp_path = None
                    temp_name = None

                    # 在 TEMP_DIR 里找对应文件：{file_id}__原文件名
                    for fn in os.listdir(TEMP_DIR):
                        if fn.startswith(prefix):
                            temp_path = os.path.join(TEMP_DIR, fn)
                            temp_name = fn.split("__", 1)[1] if "__" in fn else fn
                            break

                    if not temp_path:
                        print(f"[chat] file_id={fid} not found in {TEMP_DIR}")
                        continue

                    text = _extract_temp_text(temp_path)
                    print(f"[chat] found file={temp_name}, extracted_len={len(text)}")

                    if not text.strip():
                        continue

                    # 限长，避免 prompt 过大
                    text = text[:12000]
                    file_blocks.append(f"【附件：{temp_name}】\n{text}")

            except Exception as e:
                print(f"读取临时文件失败: {e}")

            file_context = "\n\n".join(file_blocks).strip()
        


        # === 分支 A: 专业模式 (RAG + KG) ===
        if request.mode == "course_graph":
            if rag_engine:
                yield json.dumps({
                    "type": "markdown_chunk",
                    "content": "🔍 **正在检索知识库...**\n\n"
                }, ensure_ascii=False) + "\n"
         
                # ========= KG（Excel -> 三元组） =========
                kg_context = ""
                if kg_query_tool:
                    try:
                        kg_context = kg_query_tool(user_prompt) or ""
                       
                        if "未能识别专业术语" in kg_context or kg_context.startswith("Error:"):
                            kg_context = ""
                    except Exception as e:
                        print(f"KG 查询失败: {e}")
                        kg_context = ""
                
         
                # ========= 把“附件内容 + KG结果”拼成 extra_context =========
                extra_parts = []
                if file_context:
                    extra_parts.append("【用户附件内容】\n" + file_context[:12000])
                if kg_context:
                    extra_parts.append("【知识图谱三元组】\n" + kg_context[:8000])
                extra_context = "\n\n".join(extra_parts).strip()
               
         
                try:
                    for piece in rag_engine.chat_stream(
                        user_prompt, 
                        history=request.history,
                        top_k=3, 
                        extra_context=extra_context
                    ):
                        
                        yield json.dumps({
                            "type": "markdown_chunk",
                            "content": piece,
                        }, ensure_ascii=False) + "\n"
                except Exception as e:
                    yield json.dumps({
                        "type": "error",
                        "content": f"RAG/KG 执行出错: {str(e)}"
                    }, ensure_ascii=False) + "\n"
            else:
                yield json.dumps({
                    "type": "markdown_chunk",
                    "content": "⚠️ **知识库未加载**\n\n后端未检测到索引，请先启用知识库文件并构建索引。"
                }, ensure_ascii=False) + "\n"
         

        # === 分支 B: 通用模式 (原生 DeepSeek) ===
        else:
            # 0. 检查工具 (保留你原来的逻辑)
            if request.selected_tools:
                print(f"🔧 用户激活了工具: {request.selected_tools}")

            # 1. 发送思考提示
            yield json.dumps({
                "type": "markdown_chunk",
                "content": "😎 **开始思考 **\n\n"
            }, ensure_ascii=False) + "\n"

            

            # 2. 构造消息（保留你的 Prompt）
            messages = [
                {
                    "role": "system",
                    "content": (
                        "你是一名专业的学习助手，熟悉中外物理学家生平和贡献。"
                        "用户问你任何与物理有关的问题，都应当提供结构清晰、可靠、自然的回答。"
                    )
                }
            ]

            # ✅循环注入历史记录
            for msg in (request.history or []):
                messages.append({
                    "role": msg.get("role"),
                    "content": msg.get("content")
                })

            # ✅ 注入附件内容到 system 
            if file_context:
                messages.append({
                    "role": "system",
                    "content": (
                        "下面是用户上传的附件内容（可能被截断）。回答时必须结合附件信息；"
                        "如果附件信息不足以支持结论，请明确指出缺少什么。\n\n"
                        f"{file_context}"
                    )
                })
            # ========= ✅ 注入结束 =========

            # user 消息最后放
            messages.append({"role": "user", "content": user_prompt})

            try:
                stream = await client.chat.completions.create(
                    model="deepseek-chat",
                    messages=messages,
                    stream=True,
                )

                async for chunk in stream:
                    piece = chunk.choices[0].delta.content or ""
                    if not piece:
                        continue

                    yield json.dumps({
                        "type": "markdown_chunk",
                        "content": piece,
                    }, ensure_ascii=False) + "\n"

            except Exception as e:
                yield json.dumps({
                    "type": "error",
                    "content": f"DeepSeek API 错误：{e}",
                }, ensure_ascii=False) + "\n"

        yield json.dumps({"type": "stream_end"}, ensure_ascii=False) + "\n"

    return StreamingResponse(event_generator(), media_type="text/event-stream")




if __name__ == "__main__":
    uvicorn.run("backend_main:app", host="0.0.0.0", port=8000, reload=True)